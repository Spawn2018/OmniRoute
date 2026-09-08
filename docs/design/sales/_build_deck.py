"""Buduje deck spotkania z tokenów makiety. Nie produkt, nie liczby z PLAN."""
from __future__ import annotations

import json
from base64 import b64decode
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SCREENS = ROOT / "screens"
SALES = ROOT / "sales"
LOGS = Path(r"C:\Users\sebas\.cursor\browser-logs")

INK = RGBColor(0x1A, 0x33, 0x2E)
ON_INK = RGBColor(0xF4, 0xF8, 0xF6)
CANVAS = RGBColor(0xE8, 0xF0, 0xEC)
SURFACE = RGBColor(0xF7, 0xFA, 0xF8)
TEXT = RGBColor(0x1E, 0x2E, 0x2A)
TEXT2 = RGBColor(0x4A, 0x5C, 0x57)
ACCENT = RGBColor(0x2A, 0x6B, 0x58)
ALERT = RGBColor(0x8A, 0x6A, 0x28)
STOP = RGBColor(0x9A, 0x3A, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SHOTS = {
    "home": LOGS / "cdp-response-Page.captureScreenshot-2026-09-08T00-25-17-934Z.json",
    "quote": LOGS / "cdp-response-Page.captureScreenshot-2026-09-08T00-25-24-136Z.json",
    "hitl": LOGS / "cdp-response-Page.captureScreenshot-2026-09-08T00-25-29-508Z.json",
    "ship": LOGS / "cdp-response-Page.captureScreenshot-2026-09-08T00-25-37-310Z.json",
}


def decode_shots() -> dict[str, Path]:
    out: dict[str, Path] = {}
    for name, src in SHOTS.items():
        dest = SCREENS / f"demo-{name}.png"
        if src.exists():
            payload = json.loads(src.read_text(encoding="utf-8"))
            dest.write_bytes(b64decode(payload["data"]))
            out[name] = dest
        elif dest.exists():
            out[name] = dest
        else:
            fallback = {
                "home": SCREENS / "omniroute-oferta-light.png",
                "quote": SCREENS / "omniroute-oferta-light.png",
                "hitl": SCREENS / "omniroute-hitl-light.png",
                "ship": SCREENS / "omniroute-zlecenie-light.png",
            }[name]
            out[name] = fallback
    return out


def set_run(run, size=16, bold=False, color=TEXT, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_box(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def add_text(slide, l, t, w, h, text, size=16, bold=False, color=TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, l, t, w, h, items, size=16, color=TEXT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = item
        set_run(run, size=size, color=color)
    return box


def footer(slide, n, total=11):
    add_text(slide, Inches(0.6), Inches(7.15), Inches(8), Inches(0.28),
             "OmniRoute  ·  spotkanie produktowe  ·  makieta, nie produkcja",
             size=11, color=TEXT2)
    add_text(slide, Inches(11.4), Inches(7.15), Inches(1.2), Inches(0.28),
             f"{n} / {total}", size=11, color=TEXT2, align=PP_ALIGN.RIGHT)


def shot_slide(prs, img, kicker, title, note, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(slide, 0, 0, prs.slide_width, prs.slide_height, CANVAS)
    add_box(slide, 0, 0, prs.slide_width, Inches(0.12), INK)
    add_text(slide, Inches(0.55), Inches(0.22), Inches(12), Inches(0.28), kicker, 12, False, ACCENT)
    add_text(slide, Inches(0.55), Inches(0.46), Inches(12), Inches(0.42), title, 24, True, INK)
    slide.shapes.add_picture(str(img), Inches(0.7), Inches(1.12), width=Inches(9.0))
    add_text(slide, Inches(0.7), Inches(6.78), Inches(12.0), Inches(0.32), note, 13, False, TEXT2)
    footer(slide, n)
    return slide


def build():
    shots = decode_shots()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 title / problem
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, INK)
    add_box(s, 0, 0, Inches(0.18), prs.slide_height, ACCENT)
    add_text(s, Inches(0.8), Inches(1.4), Inches(11), Inches(0.35),
             "OMNIROUTE  ·  SPEDYCJA W JEDNYM TENANCIE", 14, False, ACCENT)
    add_text(s, Inches(0.8), Inches(1.85), Inches(11.4), Inches(1.5),
             "Operator nie przepisuje maila.\nDecyduje na jednym ekranie.", 36, True, ON_INK)
    add_text(s, Inches(0.8), Inches(4.0), Inches(11), Inches(1.4),
             "Dziś praca siedzi w skrzynce: zła strona kontaktu, brak HS blokuje ofertę,\n"
             "marża ginie między arkuszem a fakturą. OmniRoute składa HITL → SQL → charge\n"
             "→ zlecenie. Model nie liczy. Człowiek zatwierdza zanim coś wejdzie do SoR.",
             18, False, ON_INK)
    add_text(s, Inches(0.8), Inches(6.6), Inches(11), Inches(0.3),
             "Spotkanie produktowe  ·  makieta działająca, nie slajd-wizja", 13, False, RGBColor(0xA8, 0xC4, 0xBA))

    # 2 status quo
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, CANVAS)
    add_box(s, 0, 0, prs.slide_width, Inches(0.12), INK)
    add_text(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.3), "KOSZT BEZCZYNNOŚCI", 12, False, ACCENT)
    add_text(s, Inches(0.7), Inches(0.7), Inches(12), Inches(0.55),
             "Rynek już zapłacił za przepisywanie i rok wdrożenia.", 28, True, INK)
    add_bullets(s, Inches(0.7), Inches(1.6), Inches(12), Inches(4.6), [
        "Podwójne wpisywanie i „umierający” TMS — powód migracji nr 1 (zestawienie case’ów Qargo / SPEED).",
        "Odpowiedzi agentów lądują w Gmailu i nie wracają do SoR — luka potwierdzona przy Magaya / overlayach.",
        "Wdrożenia legacy bolą: poślizg mierzony w miesiącach, nie w sprintach (SPEED, case’y operatorów).",
        "Rynek obiecuje mniej administracji przy AI order entry — to oczekiwanie klienta, nie nasz KPI.",
        "Czego wciąż brak: łańcuch skutków na zleceniu, HITL przed zapisem, jedna prawda o marży.",
    ], 18, TEXT)
    footer(s, 2)

    # 3 model
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, CANVAS)
    add_box(s, 0, 0, prs.slide_width, Inches(0.12), INK)
    add_text(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.3), "MODEL", 12, False, ACCENT)
    add_text(s, Inches(0.7), Inches(0.7), Inches(12), Inches(0.5),
             "Cztery zdania. Reszta to egzekucja.", 28, True, INK)
    cards = [
        ("Tenant", "Każdy wiersz ma organization_id. RLS w bazie. Zero zapytań cross-tenant."),
        ("HITL", "Model wyciąga. Kod przetwarza. Nic z ekstrakcji nie wchodzi bez akceptacji."),
        ("charge", "Jedyny wiersz kupno + sprzedaż. Marża nie mieszka w arkuszu ani w LLM."),
        ("SQL, nie model", "Wycena i matching w Postgresie. Kwota to Decimal ze stringiem na UI."),
    ]
    for i, (h, b) in enumerate(cards):
        x = Inches(0.7 + (i % 2) * 6.2)
        y = Inches(1.55 + (i // 2) * 2.35)
        add_box(s, x, y, Inches(5.9), Inches(2.15), SURFACE)
        add_text(s, x + Inches(0.25), y + Inches(0.2), Inches(5.4), Inches(0.4), h, 18, True, INK)
        add_text(s, x + Inches(0.25), y + Inches(0.7), Inches(5.4), Inches(1.2), b, 15, False, TEXT)
    footer(s, 3)

    shot_slide(prs, shots["home"], "DOWÓD NA EKRANIE  ·  DZIEŃ",
               "Joby do zamknięcia — nie skrzynka do przeszukiwania.",
               "HITL, brak HS, luka ISPS, lookup tax_id. Tenant z paska.", 4)
    shot_slide(prs, shots["hitl"], "DOWÓD  ·  EKSTRAKCJA",
               "Szkic z PDF. Accept to jedna transakcja.",
               "ExtractionService nie importuje stawek. source_ref zostaje. Art. 50 na recenzji.", 5)
    shot_slide(prs, shots["quote"], "DOWÓD  ·  WYCENA",
               "Warianty, HS, Incoterm po negocjacji. Send dopiero gdy komplet.",
               "FOB → DAP Lima zapisane jako zmiana, nie milcząca sprzeczność. Marża z charge.", 6)
    shot_slide(prs, shots["ship"], "DOWÓD  ·  ZLECENIE",
               "Legi z rolami, semafor free-time, pakiet celny.",
               "Agent ≠ armator. Dokumenty i odprawa wg Incoterms. Mapa AIS tylko tu.", 7)

    # 8 roadmap
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, CANVAS)
    add_box(s, 0, 0, prs.slide_width, Inches(0.12), INK)
    add_text(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.3), "DOKĄD ZMIERZA  ·  OŚ 2026-09-08c", 12, False, ACCENT)
    add_text(s, Inches(0.7), Inches(0.7), Inches(12), Inches(0.5),
             "Zbudowane vs w kolejce. Bez teatru.", 28, True, INK)
    add_text(s, Inches(0.7), Inches(1.45), Inches(5.8), Inches(0.35), "DZIAŁA W PRODUKCIE / MAKIECIE", 13, True, ACCENT)
    add_bullets(s, Inches(0.7), Inches(1.85), Inches(5.8), Inches(4.4), [
        "Tenant, RLS, OpenFGA, HITL accept → rate_line",
        "charge buy+sell, silnik wyceny SQL",
        "Zlecenie z rolami na legu i dokumentami",
        "Semafor storage / detention z dni serwera",
        "Pakiet celny: odbiorca × Incoterms, HITL send",
    ], 16)
    add_text(s, Inches(6.9), Inches(1.45), Inches(5.8), Inches(0.35), "W KOLEJCE PLANU (NIE OBIECUJEMY DZIŚ)", 13, True, ALERT)
    add_bullets(s, Inches(6.9), Inches(1.85), Inches(5.8), Inches(4.4), [
        "O — buy-desk na ofercie (wątki, nie drugi Outlook)",
        "I — katalog Incoterms jako dane, nie zgadywanie",
        "C — odprawa jako ścieżka, nie „odprawiamy świat”",
        "V — wieża skutków, nie druga mapa",
        "Demo-1 — pierwszy tenant na żywych danych",
    ], 16)
    footer(s, 8)

    # 9 security
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, CANVAS)
    add_box(s, 0, 0, prs.slide_width, Inches(0.12), INK)
    add_text(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.3), "BEZPIECZEŃSTWO", 12, False, ACCENT)
    add_text(s, Inches(0.7), Inches(0.7), Inches(12), Inches(0.5),
             "To nie dodatek. To warunek sprzedaży.", 28, True, INK)
    add_bullets(s, Inches(0.7), Inches(1.6), Inches(12), Inches(4.8), [
        "Izolacja tenantów w Postgres (RLS). Brak zapytania cross-tenant = błąd blokujący.",
        "Sekrety kanałów szyfrowane kluczem tenanta. Nigdy w kodzie, logu ani promptcie.",
        "Endpoint bez jawnej relacji OpenFGA = odmowa. Brak uprawnienia nie jest tostem.",
        "Zero AI na umowach i taryfach jako źródle prawdy (CI9). Model nie liczy marży, VAT, kursu.",
        "Guard na wejściu ekstrakcji. Zapis do SoR tylko po człowieku.",
    ], 18)
    footer(s, 9)

    # 10 next
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, INK)
    add_box(s, 0, 0, Inches(0.18), prs.slide_height, ACCENT)
    add_text(s, Inches(0.8), Inches(1.5), Inches(11), Inches(0.3), "NASTĘPNY KROK", 14, False, ACCENT)
    add_text(s, Inches(0.8), Inches(1.95), Inches(11.4), Inches(1.1),
             "Pilot na jednym jobie.\nNie na całym katalogu naraz.", 32, True, ON_INK)
    add_bullets(s, Inches(0.8), Inches(3.5), Inches(11), Inches(2.8), [
        "Zakres: skrzynka → HITL → wycena DAP/FOB → zlecenie z pakietem celnym.",
        "Dane: tenant klienta, stawki z PDF, zero live HTTP do armatora w pierwszym tygodniu.",
        "Sukces spotkania: uzgodniony job i kto zatwierdza Accept — nie podpis na 70 modułów.",
    ], 18, ON_INK)

    # 11 annex
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, CANVAS)
    add_box(s, 0, 0, prs.slide_width, Inches(0.12), INK)
    add_text(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.3), "ANEX  ·  RYNEK, NIE NASI KLIENCI", 12, False, ACCENT)
    add_text(s, Inches(0.7), Inches(0.7), Inches(12), Inches(0.5),
             "Ból, który już jest w case’ach TMS.", 26, True, INK)
    add_bullets(s, Inches(0.7), Inches(1.5), Inches(12), Inches(5.0), [
        "Qargo / SPEED: klienci chcą mniej przepisywania i szybszy czas-do-wartości — nie kolejny rok integracji.",
        "Magaya + overlay (Keelway): fracht w SoR, mail poza nim. Nasza przewaga: buy-desk na ofercie, HITL przed stawką.",
        "CargoWise: role na jobie działają; 12 poziomów menu nie. Bierzemy role, odrzucamy mega-menu.",
        "Self-service statusu = mniej telefonów. Portal jest w horyzoncie, nie w pierwszym pilocie.",
        "Nie cytujemy liczb OmniRoute, których nie mamy. Nie ma tu 8 minut, 15 tysięcy użytkowników ani Bayer.",
    ], 17)
    footer(s, 11)

    dest = SALES / "omniroute-spotkanie.pptx"
    prs.save(dest)
    print(dest)


if __name__ == "__main__":
    build()
