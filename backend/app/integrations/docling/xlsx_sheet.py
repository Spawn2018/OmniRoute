"""Arkusz OOXML → tekst przez openpyxl. Kwoty zostają tekstem z komórki."""

from io import BytesIO
from itertools import zip_longest
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook

from app.domain.errors import UnparseableDocument
from app.integrations.docling.parser import DocumentText


def _workbook_has_sheet_xml(raw_bytes: bytes) -> bool:
    try:
        names = ZipFile(BytesIO(raw_bytes)).namelist()
    except BadZipFile:
        return False
    return "xl/workbook.xml" in names


def _load_book(raw_bytes: bytes, *, data_only: bool) -> Workbook:
    try:
        return load_workbook(
            BytesIO(raw_bytes),
            data_only=data_only,
            read_only=True,
        )
    except (InvalidFileException, OSError, KeyError) as exc:
        raise UnparseableDocument("xlsx nieczytelne") from exc


def _pick_sheet(
    book: Workbook,
    *,
    sheet_index: int,
    sheet_name: str | None,
) -> Any:
    if sheet_name is not None:
        if sheet_name not in book.sheetnames:
            raise UnparseableDocument("sheet_name poza zakresem")
        return book[sheet_name]
    if sheet_index >= len(book.worksheets):
        raise UnparseableDocument("sheet_index poza zakresem")
    return book.worksheets[sheet_index]


def _cell_text(cached: object, formula: object) -> str | None:
    if cached is not None and str(cached).strip() != "":
        return str(cached).strip()
    if formula is not None and str(formula).strip() != "":
        # Bez cache Excel: zostaw tekst formuły; nie wyliczaj.
        return str(formula).strip()
    return None


def _sheet_lines(values_sheet: Any, formulas_sheet: Any) -> list[str]:
    lines: list[str] = []
    for value_row, formula_row in zip_longest(
        values_sheet.iter_rows(values_only=True),
        formulas_sheet.iter_rows(values_only=True),
        fillvalue=(),
    ):
        parts: list[str] = []
        for cached, formula in zip_longest(value_row, formula_row, fillvalue=None):
            text = _cell_text(cached, formula)
            if text is not None:
                parts.append(text)
        if parts:
            lines.append(" ".join(parts))
    return lines


class XlsxSheetParser:
    def parse(
        self,
        *,
        source_ref: str,
        raw_bytes: bytes,
        sheet_index: int = 0,
        sheet_name: str | None = None,
    ) -> DocumentText:
        del source_ref
        if sheet_index < 0:
            raise UnparseableDocument("sheet_index poza zakresem")
        if not _workbook_has_sheet_xml(raw_bytes):
            raise UnparseableDocument("xlsx bez skoroszytu")
        values = _load_book(raw_bytes, data_only=True)
        formulas = _load_book(raw_bytes, data_only=False)
        values_sheet = _pick_sheet(
            values,
            sheet_index=sheet_index,
            sheet_name=sheet_name,
        )
        formulas_sheet = _pick_sheet(
            formulas,
            sheet_index=sheet_index,
            sheet_name=sheet_name,
        )
        text = "\n".join(_sheet_lines(values_sheet, formulas_sheet)).strip()
        if not text:
            raise UnparseableDocument("xlsx bez tekstu")
        return DocumentText(text=text, parser_name="xlsx_sheet")
