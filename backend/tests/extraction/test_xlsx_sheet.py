from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from app.domain.errors import UnparseableDocument
from app.integrations.docling.parser import DeterministicDocumentParser, layout_fingerprint
from app.integrations.docling.xlsx_sheet import XlsxSheetParser


def _xlsx_bytes(*, text: str = "THC 10 EUR", empty: bool = False) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Sheet1"
    if not empty:
        sheet["A1"] = text
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _two_sheet_xlsx() -> bytes:
    book = Workbook()
    first = book.active
    assert first is not None
    first.title = "Sheet1"
    first["A1"] = "BAF 12 USD"
    second = book.create_sheet("Sheet2")
    second["A1"] = "THC 10 EUR"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _zip_without_workbook() -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("readme.txt", "no sheet")
    return buffer.getvalue()


def test_fingerprint_xlsx_vs_plain_zip() -> None:
    assert layout_fingerprint(_xlsx_bytes()) == "xlsx"
    assert layout_fingerprint(_zip_without_workbook()) == "text"
    assert layout_fingerprint(b"THC 10 EUR") == "text"


def test_xlsx_sheet_extracts_charge_line() -> None:
    parsed = DeterministicDocumentParser().parse(source_ref="doc://x", raw_bytes=_xlsx_bytes())
    assert parsed.parser_name == "xlsx_sheet"
    assert "THC 10 EUR" in parsed.text


def test_xlsx_empty_sheet_is_unparseable() -> None:
    with pytest.raises(UnparseableDocument, match="tekstu"):
        XlsxSheetParser().parse(source_ref="doc://x", raw_bytes=_xlsx_bytes(empty=True))


def test_xlsx_parser_requires_workbook() -> None:
    with pytest.raises(UnparseableDocument, match="skoroszytu"):
        XlsxSheetParser().parse(source_ref="doc://x", raw_bytes=_zip_without_workbook())


def test_xlsx_sheet_index_reads_second_sheet() -> None:
    raw = _two_sheet_xlsx()
    first = XlsxSheetParser().parse(source_ref="doc://x", raw_bytes=raw, sheet_index=0)
    second = XlsxSheetParser().parse(source_ref="doc://x", raw_bytes=raw, sheet_index=1)
    assert "BAF" in first.text
    assert "THC 10 EUR" in second.text


def test_xlsx_sheet_name_reads_named_sheet() -> None:
    raw = _two_sheet_xlsx()
    second = XlsxSheetParser().parse(source_ref="doc://x", raw_bytes=raw, sheet_name="Sheet2")
    assert "THC 10 EUR" in second.text


def test_xlsx_sheet_name_unknown() -> None:
    raw = _two_sheet_xlsx()
    with pytest.raises(UnparseableDocument, match="sheet_name"):
        XlsxSheetParser().parse(source_ref="doc://x", raw_bytes=raw, sheet_name="missing")


def test_xlsx_openpyxl_reads_split_cells() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "THC"
    sheet["B1"] = 10
    sheet["C1"] = "EUR"
    buffer = BytesIO()
    book.save(buffer)
    parsed = XlsxSheetParser().parse(source_ref="doc://x", raw_bytes=buffer.getvalue())
    assert parsed.parser_name == "xlsx_sheet"
    assert "THC" in parsed.text
    assert "10" in parsed.text
    assert "EUR" in parsed.text


def test_xlsx_formula_without_cache_keeps_formula_text() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "THC"
    sheet["B1"] = "=10+0"
    sheet["C1"] = "EUR"
    buffer = BytesIO()
    book.save(buffer)
    parsed = XlsxSheetParser().parse(source_ref="doc://x", raw_bytes=buffer.getvalue())
    assert "THC" in parsed.text
    assert "EUR" in parsed.text
    assert "=10+0" in parsed.text
    assert "10.0" not in parsed.text


def test_fixture_two_xlsx_still_readable_via_openpyxl() -> None:
    path = Path(__file__).resolve().parent / "fixtures" / "two.xlsx"
    second = XlsxSheetParser().parse(
        source_ref="doc://x",
        raw_bytes=path.read_bytes(),
        sheet_name="Sheet2",
    )
    assert "THC 10 EUR" in second.text
